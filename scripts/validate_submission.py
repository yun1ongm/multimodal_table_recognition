#!/usr/bin/env python3
"""Validate a submission workbook against tests.xlsx."""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from dataset_utils import as_text, load_test_questions, utc_now, write_json


NUMBER_RE = re.compile(
    r"^[+-]?(?:(?:\d+(?:\.\d*)?)|(?:\.\d+))(?:[eE][+-]?\d+)?%?$"
)


def sort_id(value: str) -> tuple[int, int | str]:
    return (0, int(value)) if value.isdigit() else (1, value)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("submission", type=Path)
    parser.add_argument("--tests", type=Path, default=Path("tests.xlsx"))
    parser.add_argument("--report", type=Path)
    parser.add_argument("--strict", action="store_true", help="Treat warnings as errors")
    return parser.parse_args()


def load_submission(path: Path) -> tuple[tuple[str, ...], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if not workbook.worksheets:
        raise ValueError("Submission workbook has no worksheets")
    worksheet = workbook.worksheets[0]
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = tuple(as_text(value) for value in header_row)
    index = {name: position for position, name in enumerate(headers)}
    if "id" not in index or "answer" not in index:
        raise ValueError(f"Submission must contain id and answer columns; got {headers}")

    rows: list[dict[str, Any]] = []
    for excel_row, row_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(row_values)
        row_id = as_text(values[index["id"]] if index["id"] < len(values) else "")
        answer = as_text(values[index["answer"]] if index["answer"] < len(values) else "")
        if not row_id and not answer and all(value is None for value in values):
            continue
        rows.append({"excel_row": excel_row, "id": row_id, "answer": answer})
    return headers, rows


def parse_json(value: str) -> tuple[Any | None, str | None]:
    if not value:
        return None, "empty answer"

    def reject_constant(token: str) -> None:
        raise ValueError(f"non-standard JSON constant {token}")

    try:
        return json.loads(value, parse_constant=reject_constant), None
    except (json.JSONDecodeError, ValueError) as exc:
        return None, str(exc)


def is_int(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def validate_structure(value: Any) -> list[str]:
    errors: list[str] = []
    if not isinstance(value, dict):
        return ["structure answer must be a JSON object"]
    for field in ("row_count", "col_count", "cells"):
        if field not in value:
            errors.append(f"missing structure field: {field}")
    if errors:
        return errors
    row_count = value["row_count"]
    col_count = value["col_count"]
    cells = value["cells"]
    if not is_int(row_count) or row_count <= 0:
        errors.append("row_count must be a positive integer")
    if not is_int(col_count) or col_count <= 0:
        errors.append("col_count must be a positive integer")
    if not isinstance(cells, list):
        return errors + ["cells must be a JSON array"]
    if errors:
        return errors

    occupied: dict[tuple[int, int], int] = {}
    required = ("text", "row", "col", "rowspan", "colspan")
    for index, cell in enumerate(cells):
        if not isinstance(cell, dict):
            errors.append(f"cells[{index}] must be an object")
            continue
        missing = [field for field in required if field not in cell]
        if missing:
            errors.append(f"cells[{index}] missing: {', '.join(missing)}")
            continue
        if not isinstance(cell["text"], str):
            errors.append(f"cells[{index}].text must be a string")
        if not all(is_int(cell[field]) for field in ("row", "col", "rowspan", "colspan")):
            errors.append(f"cells[{index}] row/col/span fields must be integers")
            continue
        if cell["row"] < 0 or cell["col"] < 0 or cell["rowspan"] <= 0 or cell["colspan"] <= 0:
            errors.append(f"cells[{index}] has invalid row/col/span value")
            continue
        row_end = cell["row"] + cell["rowspan"]
        col_end = cell["col"] + cell["colspan"]
        if row_end > row_count or col_end > col_count:
            errors.append(f"cells[{index}] exceeds row_count/col_count")
        for row in range(cell["row"], row_end):
            for col in range(cell["col"], col_end):
                key = (row, col)
                if key in occupied:
                    errors.append(
                        f"cells[{index}] overlaps cells[{occupied[key]}] at row={row}, col={col}"
                    )
                else:
                    occupied[key] = index
    return errors


def validate_answer(answer: str, answer_format: str, question_type: str) -> list[str]:
    if not answer:
        return []
    if answer_format == "number":
        if not NUMBER_RE.fullmatch(answer):
            return ["number answer is not a plain numeric value"]
        if answer.endswith("%"):
            numeric = answer[:-1]
        else:
            numeric = answer
        try:
            parsed = float(numeric)
            if not math.isfinite(parsed):
                return ["number answer is not finite"]
        except ValueError:
            return ["number answer cannot be parsed"]
        return []
    if answer_format not in {"json", "json_array"}:
        return []

    parsed, error = parse_json(answer)
    if error:
        return [f"invalid JSON: {error}"]
    if answer_format == "json_array" and not isinstance(parsed, list):
        return ["json_array answer must be a JSON array"]
    if question_type == "structure":
        return validate_structure(parsed)
    return []


def validate(tests_path: Path, submission_path: Path, strict: bool = False) -> dict[str, Any]:
    questions, _ = load_test_questions(tests_path)
    expected = {question["id"]: question for question in questions}
    headers, rows = load_submission(submission_path)

    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    if headers[:2] != ("id", "answer"):
        warnings.append({"kind": "header_order", "message": f"got {headers}"})
    if len(headers) > 2:
        warnings.append({"kind": "extra_columns", "columns": list(headers[2:])})

    actual_by_id: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        row_id = row["id"]
        if not row_id:
            errors.append({"excel_row": row["excel_row"], "kind": "blank_id"})
            continue
        actual_by_id.setdefault(row_id, []).append(row)

    duplicates = {row_id: items for row_id, items in actual_by_id.items() if len(items) > 1}
    for row_id, items in duplicates.items():
        errors.append(
            {
                "kind": "duplicate_id",
                "id": row_id,
                "excel_rows": [item["excel_row"] for item in items],
            }
        )
    missing = sorted(set(expected) - set(actual_by_id), key=sort_id)
    extra = sorted(set(actual_by_id) - set(expected), key=sort_id)
    if missing:
        errors.append({"kind": "missing_ids", "count": len(missing), "ids": missing[:50]})
    if extra:
        errors.append({"kind": "extra_ids", "count": len(extra), "ids": extra[:50]})

    blank_answers = 0
    for row_id, items in actual_by_id.items():
        if row_id not in expected or len(items) != 1:
            continue
        row = items[0]
        answer = row["answer"]
        if not answer:
            blank_answers += 1
            warnings.append({"kind": "blank_answer", "id": row_id, "excel_row": row["excel_row"]})
            continue
        question = expected[row_id]
        answer_errors = validate_answer(answer, question["answer_format"], question["question_type"])
        for message in answer_errors:
            errors.append(
                {
                    "kind": "answer_format",
                    "id": row_id,
                    "excel_row": row["excel_row"],
                    "answer_format": question["answer_format"],
                    "message": message,
                }
            )

    if strict and warnings:
        errors.extend({**warning, "promoted_from": "warning"} for warning in warnings)

    return {
        "generated_at": utc_now(),
        "tests": str(tests_path),
        "submission": str(submission_path),
        "expected_rows": len(expected),
        "actual_rows": len(rows),
        "blank_answers": blank_answers,
        "errors": errors,
        "warnings": warnings,
        "ok": not errors,
    }


def main() -> int:
    args = parse_args()
    if not args.submission.is_file():
        raise SystemExit(f"submission not found: {args.submission}")
    report = validate(args.tests.resolve(), args.submission.resolve(), strict=args.strict)
    if args.report:
        write_json(args.report.resolve(), report)
    print(
        f"submission={report['submission']} expected={report['expected_rows']} "
        f"actual={report['actual_rows']} blanks={report['blank_answers']}"
    )
    print(f"errors={len(report['errors'])} warnings={len(report['warnings'])}")
    if report["errors"]:
        for item in report["errors"][:20]:
            print("ERROR", json.dumps(item, ensure_ascii=False))
    if report["warnings"]:
        for item in report["warnings"][:5]:
            print("WARN", json.dumps(item, ensure_ascii=False))
        if len(report["warnings"]) > 5:
            print(f"WARN ... {len(report['warnings']) - 5} more warnings")
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
