"""Shared input/output utilities for the table-recognition pipeline.

This module intentionally has no Qwen dependency.  It is the deterministic
boundary around the model pipeline: question normalization, safe file
resolution, PDF metadata, and answer validation helpers.
"""

from __future__ import annotations

import json
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


EXPECTED_TEST_HEADERS = (
    "id",
    "file_name",
    "question_type",
    "question",
    "table_hint",
    "answer_format",
    "answer",
)
VALID_QUESTION_TYPES = {"structure", "extract", "thinking"}
VALID_ANSWER_FORMATS = {"string", "number", "json_array", "json"}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def as_text(value: Any) -> str:
    """Convert an Excel value into stable text without leaking ``None``."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def canonical_numeric_stem(name: str) -> str | None:
    """Return a leading-zero-insensitive numeric stem, if one exists."""

    stem = Path(name).stem.strip()
    if not stem.isdigit():
        return None
    return str(int(stem))


def infer_question_type(raw_type: str, question: str, answer_format: str) -> tuple[str, str | None]:
    """Normalize a question type and return an optional warning.

    The supplied workbook contains one invalid type (id 63).  The fallback is
    deliberately conservative: structure language plus JSON means structure;
    otherwise a valid scalar/array answer format is treated as extraction.
    """

    raw_type = raw_type.strip()
    if raw_type in VALID_QUESTION_TYPES:
        return raw_type, None

    q = question.lower()
    if answer_format == "json" and any(token in q for token in ("恢复", "结构", "表头")):
        inferred = "structure"
    else:
        inferred = "extract"
    warning = f"invalid question_type={raw_type!r}; inferred {inferred}"
    return inferred, warning


def load_test_questions(tests_path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Load and normalize the question workbook while preserving raw fields."""

    workbook = load_workbook(tests_path, read_only=True, data_only=True)
    if not workbook.worksheets:
        raise ValueError(f"Workbook has no worksheets: {tests_path}")
    worksheet = workbook.worksheets[0]
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = tuple(as_text(value) for value in header_values)
    if headers != EXPECTED_TEST_HEADERS:
        raise ValueError(
            "Unexpected tests.xlsx header. "
            f"Expected {EXPECTED_TEST_HEADERS}, got {headers}"
        )

    questions: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for excel_row, row_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        row = list(row_values)
        if len(row) < len(EXPECTED_TEST_HEADERS):
            row.extend([None] * (len(EXPECTED_TEST_HEADERS) - len(row)))
        raw = dict(zip(EXPECTED_TEST_HEADERS, row[: len(EXPECTED_TEST_HEADERS)]))
        question_id = as_text(raw["id"])
        if not question_id:
            warnings.append({"excel_row": excel_row, "kind": "blank_id"})

        question = as_text(raw["question"])
        answer_format = as_text(raw["answer_format"])
        normalized_type, type_warning = infer_question_type(
            as_text(raw["question_type"]), question, answer_format
        )
        if type_warning:
            warnings.append(
                {
                    "excel_row": excel_row,
                    "id": question_id,
                    "kind": "question_type",
                    "message": type_warning,
                }
            )
        if answer_format not in VALID_ANSWER_FORMATS:
            warnings.append(
                {
                    "excel_row": excel_row,
                    "id": question_id,
                    "kind": "answer_format",
                    "message": f"unsupported answer_format={answer_format!r}",
                }
            )

        questions.append(
            {
                "excel_row": excel_row,
                "id": question_id,
                "file_name": as_text(raw["file_name"]),
                "question_type": normalized_type,
                "raw_question_type": as_text(raw["question_type"]),
                "question": question,
                "table_hint": as_text(raw["table_hint"]),
                "answer_format": answer_format,
                "answer": as_text(raw["answer"]),
            }
        )

    return questions, {"headers": list(headers), "warnings": warnings}


def index_files(files_dir: Path) -> tuple[dict[str, Path], dict[tuple[str, str], list[Path]]]:
    """Index exact names and (extension, numeric stem) aliases."""

    exact: dict[str, Path] = {}
    aliases: dict[tuple[str, str], list[Path]] = {}
    for path in sorted(files_dir.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        exact[path.name] = path
        numeric_stem = canonical_numeric_stem(path.name)
        if numeric_stem is not None:
            key = (path.suffix.lower(), numeric_stem)
            aliases.setdefault(key, []).append(path)
    return exact, aliases


def resolve_file_name(
    requested_name: str,
    exact: dict[str, Path],
    aliases: dict[tuple[str, str], list[Path]],
) -> dict[str, Any]:
    """Resolve a workbook file reference without allowing path traversal."""

    requested_name = requested_name.strip()
    requested_path = Path(requested_name)
    if requested_path.name != requested_name or requested_name in {"", ".", ".."}:
        return {
            "requested": requested_name,
            "resolved": None,
            "status": "invalid_path",
            "candidates": [],
        }

    if requested_name in exact:
        return {
            "requested": requested_name,
            "resolved": exact[requested_name].name,
            "status": "exact",
            "candidates": [exact[requested_name].name],
        }

    numeric_stem = canonical_numeric_stem(requested_name)
    key = (requested_path.suffix.lower(), numeric_stem) if numeric_stem is not None else None
    candidates = aliases.get(key, []) if key is not None else []
    names = [path.name for path in candidates]
    if len(candidates) == 1:
        return {
            "requested": requested_name,
            "resolved": names[0],
            "status": "alias",
            "candidates": names,
        }
    if len(candidates) > 1:
        # Prefer the shortest zero-padded representation but keep the
        # ambiguity visible in the audit report.
        chosen = sorted(candidates, key=lambda path: (len(path.stem), path.name))[0]
        return {
            "requested": requested_name,
            "resolved": chosen.name,
            "status": "ambiguous_alias",
            "candidates": names,
        }
    return {
        "requested": requested_name,
        "resolved": None,
        "status": "missing",
        "candidates": [],
    }


def _run_command(command: list[str], timeout: int = 60) -> tuple[int, str, str]:
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=timeout,
            check=False,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
        return 127, "", str(exc)
    return completed.returncode, completed.stdout, completed.stderr


def pdf_metadata(path: Path) -> dict[str, Any]:
    """Collect lightweight PDF metadata using optional system utilities."""

    result: dict[str, Any] = {
        "pages": None,
        "text_chars": None,
        "text_layer": "unknown",
        "pdfinfo_available": False,
        "pdftotext_available": False,
    }
    code, stdout, _ = _run_command(["pdfinfo", str(path)])
    if code == 0:
        result["pdfinfo_available"] = True
        match = re.search(r"^Pages:\s*(\d+)\s*$", stdout, re.MULTILINE)
        if match:
            result["pages"] = int(match.group(1))

    code, stdout, _ = _run_command(["pdftotext", "-layout", str(path), "-"], timeout=120)
    if code == 0:
        result["pdftotext_available"] = True
        result["text_chars"] = len(stdout)
        result["text_layer"] = "available" if stdout.strip() else "empty"
    return result


def image_dimensions(path: Path) -> tuple[int, int] | None:
    """Read common image dimensions without requiring Pillow."""

    try:
        with path.open("rb") as handle:
            header = handle.read(32)
            if header.startswith(b"\x89PNG\r\n\x1a\n") and len(header) >= 24:
                return int.from_bytes(header[16:20], "big"), int.from_bytes(header[20:24], "big")
            if header.startswith(b"RIFF") and header[8:12] == b"WEBP":
                return _webp_dimensions(handle, header)
            if header.startswith(b"\xff\xd8"):
                return _jpeg_dimensions(handle)
    except (OSError, ValueError):
        return None
    return None


def _jpeg_dimensions(handle: Any) -> tuple[int, int] | None:
    handle.seek(2)
    while True:
        marker_prefix = handle.read(1)
        if not marker_prefix:
            return None
        if marker_prefix != b"\xff":
            continue
        marker = handle.read(1)
        while marker == b"\xff":
            marker = handle.read(1)
        if not marker:
            return None
        if marker in {b"\xd8", b"\xd9"}:
            continue
        length_bytes = handle.read(2)
        if len(length_bytes) != 2:
            return None
        length = int.from_bytes(length_bytes, "big")
        if length < 2:
            return None
        if marker[0] in set(range(0xC0, 0xC4)) | set(range(0xC5, 0xC8)) | set(range(0xC9, 0xCC)) | set(range(0xCD, 0xD0)):
            payload = handle.read(5)
            if len(payload) == 5:
                return int.from_bytes(payload[3:5], "big"), int.from_bytes(payload[1:3], "big")
            return None
        handle.seek(length - 2, 1)


def _webp_dimensions(handle: Any, header: bytes) -> tuple[int, int] | None:
    # VP8X stores 24-bit little-endian width/height minus one in its payload.
    if header[12:16] == b"VP8X":
        payload = header + handle.read(32)
        if len(payload) >= 30:
            width = 1 + int.from_bytes(payload[24:27], "little")
            height = 1 + int.from_bytes(payload[27:30], "little")
            return width, height
    return None


def json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def unique_preserving_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result
